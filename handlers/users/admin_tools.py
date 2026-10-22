import asyncio
import datetime
import re
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Any, Optional, Union

from aiogram import Bot, F, Router
from aiogram.exceptions import TelegramAPIError, TelegramForbiddenError
from aiogram.filters import BaseFilter, Command, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.types import BufferedInputFile, CallbackQuery, Message
from openpyxl import Workbook
from openpyxl.styles import Font

from filters.admin import AdminFilter
from infrastructure.database.models import User
from infrastructure.database.repo.requests import RequestsRepo
from keyboards.inline.admin import (
    broadcast_confirm_kb, broadcast_target_kb, lead_status_kb, media_list_kb, media_slot_kb,
)
from keyboards.inline.callbacks import Broadcast, BroadcastTarget, LeadStatus, MediaSlot
from misc import media
from misc.texts import REG_ROLES
from misc.media import PDF_SLOTS, PHOTO_SLOTS
from misc.states import BroadcastState, MediaUpload
from services.broadcaster import copy_to_users
from services.notify import REASON_TITLES, STATUS_TITLES, fmt_time, render_lead

MEDIA_LABELS = {**PHOTO_SLOTS, **PDF_SLOTS}

admin_tools_router = Router()
admin_tools_router.message.filter(AdminFilter())
admin_tools_router.callback_query.filter(AdminFilter())


# ---------- /stat ----------

@admin_tools_router.message(Command("stat"))
async def cmd_stat(message: Message, repo: RequestsRepo):
    users = await repo.users.stats()
    leads = await repo.leads.stats()
    lines = [
        "📊 <b>Statistika</b>",
        "",
        "👥 <b>Foydalanuvchilar</b>",
        f"Jami: {users['total']}",
        f"Ro‘yxatdan o‘tgan: {users['registered']}",
        f"Bugun qo‘shilgan: {users['today']}",
        f"Oxirgi 7 kunda: {users['week']}",
        f"Botni bloklagan: {users['blocked']}",
        "",
        "📞 <b>Arizalar (telefon raqam)</b>",
        f"Jami: {leads['total']}",
        f"Bugun: {leads['today']}",
    ]
    lines += [f"  • {REASON_TITLES.get(reason, reason)}: {count}" for reason, count in leads["by_reason"].items()]
    lines += ["", f"💡 Mavzu takliflari: {leads['suggestions']}"]
    await message.answer("\n".join(lines))


# ---------- /broadcast — hammaga yoki bitta yo'nalishga xabar yuborish ----------

async def broadcast_target_title(repo: RequestsRepo, direction_id: Optional[int]) -> str:
    if direction_id is None:
        return "👥 Hammaga"
    direction = await repo.directions.get(direction_id)
    return direction.title if direction else "?"


@admin_tools_router.message(Command("broadcast"))
async def cmd_broadcast(message: Message, state: FSMContext, repo: RequestsRepo):
    await state.set_state(None)
    await state.update_data(broadcast_message_id=None)
    directions = await repo.directions.list_all()
    counts = await repo.users.count_active_by_direction()
    total = len(await repo.users.active_ids())
    await message.answer(
        "📣 <b>Rassilka</b>\n\nKimga yuboramiz? Qavs ichida — botni bloklamagan foydalanuvchilar soni.",
        reply_markup=broadcast_target_kb(directions, counts, total),
    )


@admin_tools_router.callback_query(BroadcastTarget.filter())
async def broadcast_target(call: CallbackQuery, callback_data: BroadcastTarget, state: FSMContext, repo: RequestsRepo):
    direction_id = callback_data.direction_id or None
    await state.set_state(BroadcastState.message)
    await state.update_data(broadcast_direction_id=direction_id)
    await call.answer()
    await call.message.edit_text(
        f"📣 Kimga: <b>{await broadcast_target_title(repo, direction_id)}</b>\n\n"
        "Yubormoqchi bo‘lgan xabarni yuboring: matn, rasm, video, fayl — istalgani.\n\n"
        "Yuborishdan oldin uni sizga ko‘rsatib, tasdiqlash so‘rayman.\nBekor qilish: /cancel"
    )


@admin_tools_router.message(BroadcastState.message)
async def broadcast_preview(message: Message, state: FSMContext, bot: Bot, repo: RequestsRepo):
    await state.set_state(None)
    await state.update_data(broadcast_message_id=message.message_id)
    direction_id = (await state.get_data()).get("broadcast_direction_id")
    recipients = [uid for uid in await repo.users.active_ids(direction_id) if uid != message.from_user.id]
    await message.answer("👇 Foydalanuvchilar xabarni shunday ko‘radi:")
    await bot.copy_message(message.chat.id, message.chat.id, message.message_id)
    await message.answer(
        f"Kimga: <b>{await broadcast_target_title(repo, direction_id)}</b>\n"
        f"Shu xabar <b>{len(recipients)}</b> ta foydalanuvchiga yuboriladi. Tasdiqlaysizmi?",
        reply_markup=broadcast_confirm_kb(),
    )


async def run_broadcast(bot: Bot, session_pool, admin_id: int, message_id: int, direction_id: Optional[int]):
    """Rassilka fonda ishlaydi — bot bu vaqtda boshqa foydalanuvchilarga javob berishda davom etadi."""
    async with session_pool() as session:
        repo = RequestsRepo(session)
        user_ids = [uid for uid in await repo.users.active_ids(direction_id) if uid != admin_id]
        sent, blocked, failed = await copy_to_users(bot, user_ids, admin_id, message_id)
        await repo.users.set_inactive(blocked)
    await bot.send_message(
        admin_id,
        "✅ <b>Rassilka tugadi</b>\n\n"
        f"Yuborildi: {sent}\n"
        f"Botni bloklagan: {len(blocked)}\n"
        f"Boshqa xato: {failed}",
    )


@admin_tools_router.callback_query(Broadcast.filter(F.action == "send"))
async def broadcast_send(call: CallbackQuery, state: FSMContext, bot: Bot, session_pool):
    data = await state.get_data()
    message_id = data.get("broadcast_message_id")
    if message_id is None:
        await call.answer("Xabar topilmadi. Qaytadan /broadcast", show_alert=True)
        return
    await state.update_data(broadcast_message_id=None)
    await call.answer()
    await call.message.edit_text("⏳ Rassilka boshlandi. Tugagach hisobot yuboraman.")
    asyncio.create_task(
        run_broadcast(bot, session_pool, call.from_user.id, message_id, data.get("broadcast_direction_id"))
    )


@admin_tools_router.callback_query(Broadcast.filter(F.action == "cancel"))
async def broadcast_cancel(call: CallbackQuery, state: FSMContext):
    await state.update_data(broadcast_message_id=None)
    await call.answer()
    await call.message.edit_text("❌ Rassilka bekor qilindi.")


# ---------- Arizalar: holat, /leads, /lead_N, /export ----------

@admin_tools_router.callback_query(LeadStatus.filter())
async def change_lead_status(call: CallbackQuery, callback_data: LeadStatus, repo: RequestsRepo):
    await repo.leads.set_status(callback_data.id, callback_data.status, call.from_user.full_name)
    found = await repo.leads.get_with_user(callback_data.id)
    if found is None:
        await call.answer("Ariza topilmadi", show_alert=True)
        return
    lead, user = found
    await repo.session.refresh(user, ["direction"])
    await call.answer(f"Holat: {STATUS_TITLES[lead.status]}")
    text = render_lead(lead, user.full_name, user.username, profile=user)
    if call.message.document:
        # Plagiat arizasi — fayl, kartochka uning izohida
        await call.message.edit_caption(caption=text, reply_markup=lead_status_kb(lead.id))
    else:
        await call.message.edit_text(text, reply_markup=lead_status_kb(lead.id))


@admin_tools_router.message(Command("leads"))
async def cmd_leads(message: Message, repo: RequestsRepo):
    leads = await repo.leads.list_with_users(open_only=True, limit=30)
    if not leads:
        await message.answer("✅ Ochiq ariza yo‘q. Barcha arizalar: /export")
        return
    lines = [f"📋 <b>Ochiq arizalar</b> (yangi va bog‘lanilgan, oxirgi {len(leads)} ta)\n"]
    for lead, user in leads:
        lines.append(
            f"{STATUS_TITLES[lead.status][:1]} <b>#{lead.id}</b> · {fmt_time(lead.created_at)} · "
            f"{REASON_TITLES.get(lead.reason, lead.reason)}\n"
            f"    {escape(user.full_name)} · {escape(lead.phone or '—')} → /lead_{lead.id}"
        )
    lines.append("\nArizani ochib, holatini o‘zgartiring. Barchasi Excel'da: /export")
    await message.answer("\n".join(lines)[:4096])


@admin_tools_router.message(F.text.regexp(r"^/lead_(\d+)$").as_("match"))
async def cmd_lead(message: Message, match: re.Match, repo: RequestsRepo):
    found = await repo.leads.get_with_user(int(match.group(1)))
    if found is None:
        await message.answer("Ariza topilmadi.")
        return
    lead, user = found
    await repo.session.refresh(user, ["direction"])
    await message.answer(render_lead(lead, user.full_name, user.username, profile=user), reply_markup=lead_status_kb(lead.id))


def write_users_sheet(sheet, users: list[User]) -> None:
    sheet.append(["Ro‘yxatdan o‘tgan", "Ism-familiya", "Telefon", "Faoliyat yo‘nalishi", "Yo‘nalishi",
                  "Telegram ismi", "Username", "Telegram ID"])
    for user in users:
        sheet.append([
            fmt_time(user.registered_at), user.real_name, user.phone, REG_ROLES.get(user.role, user.role),
            user.direction.title if user.direction else "", user.full_name,
            f"@{user.username}" if user.username else "", user.user_id,
        ])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFGH", [17, 28, 16, 34, 26, 25, 18, 14]):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"


@admin_tools_router.message(Command("export"))
async def cmd_export(message: Message, repo: RequestsRepo):
    leads = await repo.leads.list_with_users()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Arizalar"
    headers = ["№", "Sana", "Bo‘lim", "Tafsilot", "Ism", "Username", "Telegram ID", "Telefon",
               "Holat", "Holatni o‘zgartirgan", "Holat sanasi"]
    sheet.append(headers)
    for lead, user in leads:
        sheet.append([
            lead.id, fmt_time(lead.created_at), REASON_TITLES.get(lead.reason, lead.reason), lead.details or "",
            user.full_name, f"@{user.username}" if user.username else "", user.user_id, lead.phone or "",
            STATUS_TITLES.get(lead.status, lead.status), lead.status_by or "", fmt_time(lead.status_at),
        ])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFGHIJK", [6, 17, 34, 40, 25, 18, 14, 16, 16, 22, 17]):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    users = await repo.users.registered()
    write_users_sheet(workbook.create_sheet("Foydalanuvchilar"), users)
    for direction in await repo.directions.list_all():
        # Excel varag'i nomi 31 belgidan oshmasligi kerak
        write_users_sheet(
            workbook.create_sheet(direction.name[:31]),
            [user for user in users if user.direction_id == direction.id],
        )

    buffer = BytesIO()
    workbook.save(buffer)
    filename = f"arizalar_{datetime.date.today():%Y-%m-%d}.xlsx"
    await message.answer_document(
        BufferedInputFile(buffer.getvalue(), filename=filename),
        caption=f"📊 Jami arizalar: {len(leads)}\n👤 Ro‘yxatdan o‘tganlar: {len(users)}",
    )


# ---------- /media — rasm va PDF'larni yuklash ----------

MEDIA_LIST_TEXT = (
    "🖼 <b>Rasm va PDF'lar</b>\n\n"
    "✅ — yuklangan, ➖ — yo‘q (bu joyda rasm yoki fayl yuborilmaydi).\n"
    "Almashtirish yoki yuklash uchun tanlang:"
)


@admin_tools_router.message(Command("media"))
async def cmd_media(message: Message):
    await message.answer(MEDIA_LIST_TEXT, reply_markup=media_list_kb())


@admin_tools_router.callback_query(MediaSlot.filter(F.action == "list"))
async def media_list(call: CallbackQuery):
    await call.answer()
    await call.message.edit_text(MEDIA_LIST_TEXT, reply_markup=media_list_kb())


@admin_tools_router.callback_query(MediaSlot.filter(F.action == "open"))
async def media_open(call: CallbackQuery, callback_data: MediaSlot, state: FSMContext):
    await call.answer()
    key = callback_data.key
    await state.set_state(MediaUpload.file)
    await state.update_data(media_key=key)
    if media.find(key):
        await call.message.answer("Hozirgisi:")
        await (media.send_document if key in PDF_SLOTS else media.send_photo)(call.message, key)
    what = "PDF fayl" if key in PDF_SLOTS else "rasm (oddiy rasm yoki .jpg/.png fayl)"
    await call.message.answer(
        f"«{MEDIA_LABELS[key]}» uchun yangi {what} yuboring.\nBekor qilish: /cancel",
        reply_markup=media_slot_kb(key),
    )


@admin_tools_router.callback_query(MediaSlot.filter(F.action == "delete"))
async def media_delete(call: CallbackQuery, callback_data: MediaSlot, state: FSMContext):
    media.delete(callback_data.key)
    await state.set_state(None)
    await call.answer("O‘chirildi")
    await call.message.edit_text(MEDIA_LIST_TEXT, reply_markup=media_list_kb())


@admin_tools_router.message(MediaUpload.file)
async def media_upload(message: Message, state: FSMContext, bot: Bot):
    key = (await state.get_data()).get("media_key")
    document = message.document
    if key in PDF_SLOTS:
        is_pdf = document and (document.mime_type == "application/pdf" or (document.file_name or "").lower().endswith(".pdf"))
        if not is_pdf:
            await message.answer("Bu yerga PDF fayl kerak. Qaytadan yuboring yoki /cancel")
            return
        await media.save_upload(bot, key, document.file_id, document.file_name, ".pdf")
    elif message.photo:
        await media.save_upload(bot, key, message.photo[-1].file_id, None, ".jpg")
    elif document and (document.mime_type or "").startswith("image/"):
        extension = Path(document.file_name or "").suffix.lower() or ".jpg"
        await media.save_upload(bot, key, document.file_id, None, extension, reuse_file_id=False)
    else:
        await message.answer("Bu yerga rasm kerak. Qaytadan yuboring yoki /cancel")
        return
    await state.set_state(None)
    await message.answer(f"✅ Saqlandi: {MEDIA_LABELS[key]}\n\n{MEDIA_LIST_TEXT}", reply_markup=media_list_kb())


# ---------- Foydalanuvchiga javob yozish (reply) ----------

class ReplyToUser(BaseFilter):
    """
    Admin bot yuborgan "✉️ Foydalanuvchidan xabar" / "🆕 Yangi ariza" xabariga yoki forward qilingan
    xabarning o'ziga reply qilganmi? Ha bo'lsa — handlerga target_user_id beriladi.
    """

    async def __call__(self, message: Message, bot: Bot) -> Union[bool, dict[str, Any]]:
        replied = message.reply_to_message
        if replied is None or replied.from_user is None or replied.from_user.id != bot.id:
            return False
        sender = getattr(replied.forward_origin, "sender_user", None)
        if sender is not None:
            return {"target_user_id": sender.id}
        found = re.search(r"ID: (\d+)", replied.text or replied.caption or "")
        return {"target_user_id": int(found.group(1))} if found else False


@admin_tools_router.message(StateFilter(None), ReplyToUser())
async def reply_to_user(message: Message, bot: Bot, target_user_id: int):
    try:
        await bot.copy_message(target_user_id, message.chat.id, message.message_id)
    except TelegramForbiddenError:
        await message.reply("❌ Yuborilmadi: foydalanuvchi botni bloklagan.")
    except TelegramAPIError as e:
        await message.reply(f"❌ Yuborilmadi: {e.message}")
    else:
        await message.reply("✅ Javob yuborildi.")
